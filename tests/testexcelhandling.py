from pathlib import Path

from openpyxl import Workbook, load_workbook
from playwright.sync_api import Page, expect


EXCELFILE = Path(__file__).resolve().parents[1] / "test_data" / "logindata.xlsx"


def testreadexcelcell():
    workbook = load_workbook(EXCELFILE)
    sheet = workbook["LoginData"]

    username = sheet["A2"].value
    password = sheet["B2"].value

    assert username == "meghana"
    assert password == "test123"

    print("Username:", username)
    print("Password retrieved successfully")

    workbook.close()


def testreadexcelrows():
    workbook = load_workbook(EXCELFILE)
    sheet = workbook["LoginData"]

    rows = list(sheet.iter_rows(min_row=2, max_col=3, values_only=True))

    assert len(rows) == 4

    for row in rows:
        print("Test data:", row)

    workbook.close()


def testexcelworkbookdetails():
    workbook = load_workbook(EXCELFILE)
    sheet = workbook["LoginData"]

    print("Sheets:", workbook.sheetnames)
    print("Rows:", sheet.max_row)
    print("Columns:", sheet.max_column)

    assert "LoginData" in workbook.sheetnames
    assert sheet["A1"].value == "Username"

    workbook.close()


def testexcelwritenewworkbook(tmp_path):
    outputfile = tmp_path / "newresults.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Results"

    sheet.append(["Test Case", "Expected Result", "Actual Result", "Status"])
    sheet.append(["TC001", "Success", "Success", "PASS"])

    workbook.save(outputfile)
    workbook.close()

    saved_workbook = load_workbook(outputfile)
    saved_sheet = saved_workbook["Results"]

    assert saved_sheet["A2"].value == "TC001"
    assert saved_sheet["D2"].value == "PASS"

    saved_workbook.close()


def testexcelupdatecell(tmp_path):
    outputfile = tmp_path / "updateresults.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Results"
    sheet.append(["Username", "Status"])
    sheet.append(["meghana", "Pending"])
    workbook.save(outputfile)
    workbook.close()

    workbook = load_workbook(outputfile)
    sheet = workbook["Results"]

    sheet["B2"] = "PASS"

    workbook.save(outputfile)
    workbook.close()

    workbook = load_workbook(outputfile, data_only=True)
    assert workbook["Results"]["B2"].value == "PASS"
    workbook.close()


def testexcelappendrow(tmp_path):
    outputfile = tmp_path / "appendresults.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Results"
    sheet.append(["Username", "Status"])
    sheet.append(["user1", "PASS"])
    workbook.save(outputfile)
    workbook.close()

    workbook = load_workbook(outputfile)
    sheet = workbook["Results"]

    sheet.append(["user2", "FAIL"])

    workbook.save(outputfile)
    workbook.close()

    workbook = load_workbook(outputfile, data_only=True)
    sheet = workbook["Results"]

    assert sheet.max_row == 3
    assert sheet["A3"].value == "user2"
    assert sheet["B3"].value == "FAIL"

    workbook.close()


def testexceldatawithplaywright(page: Page):
    workbook = load_workbook(EXCELFILE, data_only=True)
    sheet = workbook["LoginData"]

    username = sheet["A2"].value
    password = sheet["B2"].value
    expected_result = sheet["C2"].value

    workbook.close()

    page.set_content("""
        <label for="username">Username</label>
        <input id="username">

        <label for="password">Password</label>
        <input id="password" type="password">

        <button id="login">Login</button>
        <p id="result"></p>

        <script>
            document.querySelector("#login").addEventListener("click", () => {
                const username = document.querySelector("#username").value;
                const password = document.querySelector("#password").value;

                document.querySelector("#result").textContent =
                    username === "meghana" && password === "test123"
                    ? "Success"
                    : "Failure";
            });
        </script>
    """)

    page.locator("#username").fill(username)
    page.locator("#password").fill(password)
    page.locator("#login").click()

    expect(page.locator("#result")).to_have_text(expected_result)

    actual_result = page.locator("#result").inner_text()
    print("Expected:", expected_result)
    print("Actual:", actual_result)


def testexceldata_driven_login(page: Page):
    workbook = load_workbook(EXCELFILE, data_only=True)
    sheet = workbook["LoginData"]

    testdata = list(
        sheet.iter_rows(min_row=2, max_col=3, values_only=True)
    )

    workbook.close()

    page.set_content("""
        <input id="username">
        <input id="password" type="password">
        <button id="login">Login</button>
        <p id="result"></p>

        <script>
            document.querySelector("#login").addEventListener("click", () => {
                const username = document.querySelector("#username").value;
                const password = document.querySelector("#password").value;

                if (!username || !password) {
                    document.querySelector("#result").textContent =
                        "Validation Message";
                } else if (
                    username === "meghana" && password === "test123"
                ) {
                    document.querySelector("#result").textContent = "Success";
                } else {
                    document.querySelector("#result").textContent = "Failure";
                }
            });
        </script>
    """)

    for row_number, row in enumerate(testdata, start=2):
        username, password, expected_result = row

        page.locator("#username").fill(username or "")
        page.locator("#password").fill(password or "")
        page.locator("#login").click()

        actual_result = page.locator("#result").inner_text()
        status = "PASS" if actual_result == expected_result else "FAIL"

        workbook = load_workbook(EXCELFILE)
        sheet = workbook["LoginData"]
        sheet.cell(row=row_number, column=4, value=actual_result)
        sheet.cell(row=row_number, column=5, value=status)
        workbook.save(EXCELFILE)
        workbook.close()

        print(
            f"Row {row_number}: Expected={expected_result}, "
            f"Actual={actual_result}, Status={status}"
        )

        assert actual_result == expected_result, (
            f"Excel row {row_number} failed: "
            f"expected {expected_result!r}, got {actual_result!r}"
        )


def testexcelmissingvalues():
    workbook = load_workbook(EXCELFILE, data_only=True)
    sheet = workbook["LoginData"]

    username = sheet["A5"].value
    password = sheet["B5"].value

    assert username is None
    assert password is None

    print("Empty Excel cells handled correctly")

    workbook.close()